#!/usr/bin/env python3
"""Index original course materials and render real PDF/image evidence."""

import argparse
import hashlib
import json
from pathlib import Path
import sys


def require(condition, message):
    if not condition:
        raise ValueError(message)


def pdf_library():
    try:
        import pymupdf
        return pymupdf
    except ImportError as error:
        raise ValueError("PDF operations require PyMuPDF in the selected interpreter") from error


def index(source):
    extension = source.suffix.lower()
    units = []
    if extension == ".pdf":
        pymupdf = pdf_library()
        with pymupdf.open(source) as document:
            require(not document.needs_pass, "PDF is password-protected")
            for number, page in enumerate(document, 1):
                text = page.get_text("text")
                units.append({"unit": number, "kind": "page", "text": text,
                              "printed_label_hint": page.get_label(),
                              "text_present": bool(text.strip()), "visually_checked": False,
                              "studied": False})
    elif extension == ".pptx":
        try:
            from pptx import Presentation
        except ImportError as error:
            raise ValueError("PPTX indexing requires python-pptx") from error
        presentation = Presentation(source)
        for number, slide in enumerate(presentation.slides, 1):
            texts = []

            def extract(shapes):
                for shape in shapes:
                    if hasattr(shape, "shapes"):
                        extract(shape.shapes)
                    if shape.has_text_frame:
                        texts.append(shape.text_frame.text)
                    if shape.has_table:
                        texts.extend("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)

            extract(slide.shapes)
            notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ""
            units.append({"unit": number, "kind": "slide", "text": "\n".join(texts), "notes": notes,
                          "text_present": bool("".join(texts).strip() or notes.strip()),
                          "visually_checked": False, "studied": False})
    elif extension in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}:
        from PIL import Image
        with Image.open(source) as image:
            require(getattr(image, "n_frames", 1) == 1, "Multi-frame images require explicit frame extraction")
            units.append({"unit": 1, "kind": "image", "text": "", "text_present": False,
                          "dimensions": list(image.size), "visually_checked": False, "studied": False})
    else:
        raise ValueError("Supported index formats: PDF, PPTX, single-frame images; convert other formats with a faithful reader")
    require(bool(units), "Source contains no units")
    return {"total_units": len(units), "units": units,
            "note": "Extracted text is not verified learning; inspect original pages for formulas, diagrams, and OCR needs."}


def render(source, unit, crop, dpi, out):
    extension = source.suffix.lower()
    require(out.suffix.lower() == ".png", "Evidence output must have a .png extension")
    require(72 <= dpi <= 600, "DPI must be between 72 and 600")
    if crop:
        x0, y0, x1, y1 = crop
        require(0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1, "Crop must be a nonempty normalized rectangle")
    out.parent.mkdir(parents=True, exist_ok=True)
    if extension == ".pdf":
        pymupdf = pdf_library()
        with pymupdf.open(source) as document:
            require(not document.needs_pass, "PDF is password-protected")
            require(1 <= unit <= len(document), "PDF physical page is out of range")
            page = document[unit - 1]
            pixmap = page.get_pixmap(dpi=dpi, alpha=False)
            if crop:
                from PIL import Image
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                box = (int(crop[0] * image.width), int(crop[1] * image.height),
                       int(crop[2] * image.width), int(crop[3] * image.height))
                require(box[2] > box[0] and box[3] > box[1], "Crop is smaller than a pixel")
                image.crop(box).save(out, format="PNG")
            else:
                pixmap.save(str(out))
    elif extension in {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}:
        require(unit == 1, "A single-frame image has only unit 1")
        from PIL import Image, ImageOps
        with Image.open(source) as original:
            require(getattr(original, "n_frames", 1) == 1, "Multi-frame images require explicit frame extraction")
            image = ImageOps.exif_transpose(original).convert("RGB")
            if crop:
                box = (int(crop[0] * image.width), int(crop[1] * image.height),
                       int(crop[2] * image.width), int(crop[3] * image.height))
                require(box[2] > box[0] and box[3] > box[1], "Crop is smaller than a pixel")
                image = image.crop(box)
            image.save(out, format="PNG")
    else:
        raise ValueError("Faithful PPT screenshots require rendering the actual deck to PDF/images first")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    for command in ("index", "render"):
        part = sub.add_parser(command)
        part.add_argument("--source", required=True, type=Path)
        part.add_argument("--out", required=True, type=Path)
        if command == "render":
            part.add_argument("--unit", required=True, type=int)
            part.add_argument("--crop", nargs=4, type=float)
            part.add_argument("--dpi", type=int, default=160)
    args = parser.parse_args()
    source = args.source.expanduser().resolve()
    out = args.out.expanduser().resolve()
    require(source.is_file(), f"Source is unavailable: {source}")
    require(out != source and out.with_suffix(out.suffix + ".json") != source, "Output cannot overwrite the original source")
    with source.open("rb") as handle:
        digest = hashlib.file_digest(handle, "sha256").hexdigest()
    provenance = {"source_path": str(source), "source_sha256": digest}
    if args.command == "index":
        require(out.suffix.lower() == ".json", "Index output must have a .json extension")
        provenance.update(index(source))
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
    else:
        render(source, args.unit, args.crop, args.dpi, out)
        provenance.update({"unit": args.unit, "crop_normalized": args.crop, "dpi": args.dpi,
                           "image_path": str(out), "origin": "original_source_render"})
        out.with_suffix(out.suffix + ".json").write_text(json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(out), "source_sha256": digest}, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError, ImportError) as error:
        print(f"Error: {error}", file=sys.stderr)
        sys.exit(1)
